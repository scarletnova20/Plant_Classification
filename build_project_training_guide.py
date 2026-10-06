"""Build the project training and deep learning fundamentals guide."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "docs"
ASSET_DIR = OUT_DIR / "assets"
OUT_PATH = OUT_DIR / "Plant_Classification_PyTorch_Training_Guide.docx"

NAVY = "1F4E78"
BLUE = "DDEBF7"
PALE_BLUE = "EEF5FA"
PALE_GRAY = "F3F4F6"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_repeat_font(run, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char1, instr_text, fld_char2))
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_numbered(doc, text):
    add_numbered.counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(f"{add_numbered.counter}.  {text}")
    return p


add_numbered.counter = 0


def reset_numbering():
    add_numbered.counter = 0


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.right_indent = Cm(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GRAY)
    p_pr.append(shd)
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_repeat_font(run, "Consolas")
        run.font.size = Pt(8.2)
        run.font.color.rgb = RGBColor.from_string("202124")
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(header))
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(font_size)
        if widths:
            cell.width = Cm(widths[idx])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if idx > 0 and len(str(value)) < 20:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            if widths:
                cell.width = Cm(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def draw_flow(path):
    width, height = 1600, 980
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = Path(r"C:\Windows\Fonts\arial.ttf")
    bold_path = Path(r"C:\Windows\Fonts\arialbd.ttf")
    regular = ImageFont.truetype(str(regular_path), 29)
    small = ImageFont.truetype(str(regular_path), 23)
    bold = ImageFont.truetype(str(bold_path), 31)
    title = ImageFont.truetype(str(bold_path), 38)

    draw.text((70, 40), "Plant classification training flow", fill="#000000", font=title)
    boxes = [
        ("1  Dataset and manifest", "252 active images\n13 class folders"),
        ("2  Validation checks", "Paths, hashes, image readability\nand source metadata"),
        ("3  Training eligibility", "Exclude 10 sign-bearing originals\n242 usable images"),
        ("4  Source-group split", "194 train   48 validation\nNo source overlap"),
        ("5  Preprocess", "RGB   128 x 128\nNormalize values to 0-1"),
        ("6  Train the CNN", "Augmentation   convolution   pooling\ndense layer   dropout   logits"),
        ("7  Save and evaluate", "Best checkpoint and class map\n54.17% diagnostic validation accuracy"),
    ]
    x1, x2 = 130, 1470
    y = 125
    box_h = 95
    gap = 32
    for idx, (heading, body) in enumerate(boxes):
        fill = "#DDEBF7" if idx % 2 == 0 else "#EEF5FA"
        draw.rounded_rectangle((x1, y, x2, y + box_h), radius=18, fill=fill, outline="#1F4E78", width=3)
        draw.text((x1 + 30, y + 15), heading, fill="#000000", font=bold)
        draw.multiline_text((x1 + 530, y + 13), body, fill="#24364B", font=small, spacing=4)
        if idx < len(boxes) - 1:
            cx = width // 2
            arrow_top = y + box_h + 5
            arrow_bottom = y + box_h + gap - 4
            draw.line((cx, arrow_top, cx, arrow_bottom), fill="#1F4E78", width=5)
            draw.polygon([(cx - 12, arrow_bottom - 12), (cx + 12, arrow_bottom - 12), (cx, arrow_bottom + 5)], fill="#1F4E78")
        y += box_h + gap
    image.save(path, quality=95)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_term(doc, term, explanation):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(term + ". ")
    r.bold = True
    p.add_run(explanation)


def build_document():
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)
    flow_path = ASSET_DIR / "project_flow.png"
    draw_flow(flow_path)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 15, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "PLANT CLASSIFICATION PROJECT"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.runs[0]
    set_repeat_font(hr)
    hr.bold = True
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    # Cover page
    doc.add_paragraph().paragraph_format.space_after = Pt(42)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Plant Classification Training Guide")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("Project flow and PyTorch fundamentals with TensorFlow comparison")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(14)
    intro.add_run(
        "This guide explains how the plant image classifier moves from the dataset and manifest "
        "to a trained convolutional neural network. It also teaches the deep learning concepts "
        "used by the PyTorch implementation and shows how the same ideas are expressed in TensorFlow."
    )

    add_table(doc, ["Current project fact", "Value"], [
        ("Active dataset", "252 images across 13 classes"),
        ("Training eligible", "242 images after sign exclusion"),
        ("Split", "194 training and 48 validation"),
        ("Model", "Small CNN with 268,637 parameters"),
        ("Best checkpoint", "Epoch 23"),
        ("Diagnostic validation result", "54.17% accuracy and 1.4408 loss"),
    ], widths=[7.4, 9.6], font_size=9.2)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Prepared for the Plant Classification deep learning project")
    r.bold = True
    r.font.size = Pt(10)
    p2 = doc.add_paragraph("Current PyTorch model run: 6 October 2026")
    p2.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    doc.add_page_break()

    # Contents
    add_heading(doc, "Contents", 1)
    contents = [
        "1  Project architecture and file map",
        "2  End to end training flow",
        "3  Dataset validation and leakage control",
        "4  Image preprocessing and augmentation",
        "5  CNN architecture used by the project",
        "6  How model training works",
        "7  Evaluation and prediction",
        "8  Deep learning fundamentals",
        "9  TensorFlow and PyTorch comparison",
        "10  Current results and limitations",
        "11  Running the project",
        "12  Glossary and study path",
    ]
    for item in contents:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(5)

    add_heading(doc, "How to use this guide", 2)
    doc.add_paragraph(
        "Read sections 1 through 7 to understand the codebase. Sections 8 and 9 explain the "
        "mathematics and framework concepts behind that code. The final sections interpret the "
        "current results and provide commands and a practical study sequence."
    )

    # 1
    add_heading(doc, "1 Project architecture and file map", 1)
    doc.add_paragraph(
        "The training code is intentionally small. Data preparation scripts create and audit the "
        "dataset, while three core programs train, evaluate and use the classifier."
    )
    add_table(doc, ["File or folder", "Role in the project"], [
        ("train.py", "Validates the dataset, creates the split, preprocesses images, builds the CNN and trains it."),
        ("evaluate.py", "Loads the saved model and held-out split, then produces metrics and a confusion matrix."),
        ("predict.py", "Loads one new image and returns the most likely plant class and softmax score."),
        ("dataset_manifest.csv", "Records IDs, labels, provenance, source groups, collection groups, QC status and hashes."),
        ("Plant_Dataset", "Contains one folder per active class, with original and generated subfolders."),
        ("models", "Stores the PyTorch checkpoint, class-name mapping and exact saved split."),
        ("results", "Stores accuracy and loss curves, the classification report and confusion matrix."),
        ("audit and integration scripts", "Prepare, label and integrate images before training. They are not part of the training loop."),
    ], widths=[5.2, 11.8])

    add_heading(doc, "The three identifiers that protect the dataset", 2)
    add_term(doc, "image_id", "Uniquely identifies one file, such as an original or a generated variant.")
    add_term(doc, "source_group", "Links an original photograph to every generated image derived from it.")
    add_term(doc, "collection_group", "Links photographs from the same physical plant, site or collection session.")

    # 2
    add_heading(doc, "2 End to end training flow", 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(flow_path), width=Cm(16.6))
    cap = doc.add_paragraph("Figure 1  Current plant classification pipeline")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    cap.paragraph_format.space_after = Pt(10)

    reset_numbering()
    for text in [
        "The class folders and manifest are read together. Training stops if the files and metadata disagree.",
        "Images with readable botanical signs are excluded so the network cannot solve the task by reading labels.",
        "The usable data is divided by source group, preserving the relationship between originals and generated images.",
        "Images are converted to RGB, resized to 128 by 128 pixels and normalized to values between zero and one.",
        "The CNN learns filters from the training subset. Validation images measure learning after each epoch.",
        "The checkpoint with the lowest validation loss is saved and later evaluated on the same held-out subset.",
    ]:
        add_numbered(doc, text)

    # 3
    add_heading(doc, "3 Dataset validation and leakage control", 1)
    add_heading(doc, "What load_dataset checks", 2)
    reset_numbering()
    for text in [
        "Plant_Dataset and dataset_manifest.csv both exist.",
        "Every active manifest path points to a readable JPEG or PNG inside an active class folder.",
        "Every active file is represented in the manifest and every image ID is unique.",
        "Each class contains at least two active images.",
        "Every source group maps consistently to one collection group.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Why a random file split would be misleading", 2)
    doc.add_paragraph(
        "A generated image can be visually very close to its original. If the original were in training and the generated "
        "version were in validation, the model could appear accurate because it had already seen the same biological and "
        "photographic source. The project avoids this by assigning complete source groups to one subset."
    )
    add_table(doc, ["Split property", "Current value", "Reason"], [
        ("Requested validation fraction", "20%", "Provides a larger and more stable diagnostic sample."),
        ("Actual validation fraction", "19.8%", "Whole source groups prevent a mathematically exact split in every class."),
        ("Training images", "194", "Used to update the neural network weights."),
        ("Validation images", "48", "Used to select the best checkpoint and report diagnostic metrics."),
        ("Independent test images", "0", "The dataset lacks enough independent collection groups."),
        ("Source overlap", "0 groups", "Prevents original and generated derivatives from crossing the split."),
    ], widths=[5.0, 3.2, 8.8])

    # 4
    add_heading(doc, "4 Image preprocessing and augmentation", 1)
    add_heading(doc, "Preprocessing", 2)
    doc.add_paragraph(
        "Pillow opens each image and converts it to three-channel RGB. torchvision resizes it to 128 by 128 pixels, converts "
        "it to a 32-bit floating-point tensor and divides pixel values by 255. A black pixel becomes 0.0 and a maximum-intensity "
        "pixel becomes 1.0."
    )
    add_code(doc, "transform = transforms.Compose([\n    transforms.Resize((128, 128)),\n    transforms.ToTensor(),\n])\nimage_tensor = transform(image.convert(\"RGB\"))")

    add_heading(doc, "Tensor shapes", 2)
    add_table(doc, ["Stage", "Tensor shape", "Meaning"], [
        ("One image", "3 x 128 x 128", "RGB channels, height and width."),
        ("Training batch", "8 x 3 x 128 x 128", "Eight images processed before one optimizer update."),
        ("Model output", "8 x 13", "Thirteen class scores for each image in the batch."),
    ], widths=[4.0, 5.0, 8.0])

    add_heading(doc, "Training augmentation", 2)
    add_table(doc, ["Layer", "Setting", "Learning purpose"], [
        ("RandomHorizontalFlip", "Default probability 0.5", "Reduces sensitivity to left-right orientation."),
        ("RandomRotation", "18 degrees", "Adds mild rotation while preserving plant identity."),
        ("RandomAffine scale", "0.90 to 1.10", "Adds small scale changes."),
        ("RandomAffine translate", "0.08 each axis", "Moves the subject within the frame."),
    ], widths=[4.3, 4.0, 8.7])
    doc.add_paragraph(
        "The training DataLoader uses these random transformations. Validation and prediction use a separate deterministic "
        "transform containing only resize and tensor conversion, which keeps evaluation repeatable."
    )

    # 5
    add_heading(doc, "5 CNN architecture used by the project", 1)
    add_table(doc, ["Layer", "Output shape", "Parameters", "Purpose"], [
        ("Input and augmentation", "3 x 128 x 128", "0", "Accept and perturb training images."),
        ("Conv2d 8", "8 x 128 x 128", "224", "Learn low-level colors, edges and textures."),
        ("MaxPool", "8 x 64 x 64", "0", "Reduce spatial resolution."),
        ("Conv2d 16", "16 x 64 x 64", "1,168", "Learn more detailed local patterns."),
        ("MaxPool", "16 x 32 x 32", "0", "Reduce spatial resolution."),
        ("Conv2d 32", "32 x 32 x 32", "4,640", "Learn higher-level plant features."),
        ("MaxPool", "32 x 16 x 16", "0", "Produce a compact feature map."),
        ("Flatten", "8,192", "0", "Convert feature maps into one vector."),
        ("Dense 32", "32", "262,176", "Combine features for classification."),
        ("Dropout", "32", "0", "Randomly suppress 30% of dense activations in training."),
        ("Linear output", "13", "429", "Return one raw logit per plant class."),
        ("Total", "", "268,637", "All trainable weights and biases."),
    ], widths=[4.0, 3.6, 3.0, 6.4], font_size=8.2)

    add_heading(doc, "How a convolution detects a feature", 2)
    doc.add_paragraph(
        "A convolutional filter is a small grid of learned numbers. It slides across an image and produces a strong response "
        "when the local pixels match the pattern encoded by that filter. Early filters may respond to green-to-background "
        "edges. Later filters combine earlier responses into leaf shapes, veins, flower clusters or rosette structures."
    )
    doc.add_paragraph(
        "For the first convolution, each filter contains 3 x 3 x 3 weights plus one bias. With eight filters, the parameter "
        "count is (3 x 3 x 3 + 1) x 8 = 224."
    )

    # 6
    add_heading(doc, "6 How model training works", 1)
    add_heading(doc, "The learning cycle", 2)
    reset_numbering()
    for text in [
        "Forward pass: the current model converts an image batch into raw class logits.",
        "Loss calculation: CrossEntropyLoss measures how strongly the logits support the correct class.",
        "Backpropagation: PyTorch autograd calculates how each trainable parameter contributed to the loss.",
        "Optimizer update: Adam changes the parameters in the direction expected to reduce future loss.",
        "Validation: after the epoch, the model predicts the held-out images without updating parameters.",
    ]:
        add_numbered(doc, text)

    add_heading(doc, "Important training terms", 2)
    add_term(doc, "Epoch", "One complete pass through all 194 training images.")
    add_term(doc, "Batch", "A small group of images processed together. This project uses eight images per batch.")
    add_term(doc, "Learning rate", "The scale of each optimizer update. The project uses 0.001.")
    add_term(doc, "Loss", "A differentiable error value used to train the model. Lower validation loss is preferred.")
    add_term(doc, "Accuracy", "The fraction of predictions whose highest-scoring class is correct.")
    add_term(doc, "Gradient", "The direction and strength of change in loss with respect to a parameter.")

    add_heading(doc, "Cross entropy loss", 2)
    doc.add_paragraph(
        "PyTorch CrossEntropyLoss combines log-softmax with negative log-likelihood. For one image, the result is approximately "
        "the negative logarithm of the probability assigned to the correct class. A correct-class probability of 0.80 gives a "
        "small loss of about 0.22, while a probability of 0.10 gives a much larger loss of about 2.30."
    )

    add_heading(doc, "Adam, early stopping and checkpointing", 2)
    doc.add_paragraph(
        "Adam maintains moving estimates of recent gradients and squared gradients, allowing different parameters to receive "
        "adapted update sizes. Early stopping watches validation loss and can stop after five epochs without improvement. "
        "The training loop saves a checkpoint whenever validation loss improves, so the final file is not simply the last epoch."
    )
    add_code(doc, "criterion = nn.CrossEntropyLoss()\noptimizer = torch.optim.Adam(model.parameters(), lr=0.001)\nbatch_size = 8\nmaximum_epochs = 25\nearly_stopping_patience = 5")

    # 7
    add_heading(doc, "7 Evaluation and prediction", 1)
    add_heading(doc, "Evaluation output", 2)
    doc.add_paragraph(
        "evaluate.py reloads the exact saved split and checks that the class-folder ordering still matches class_names.json. It "
        "then computes loss, accuracy, per-class precision, recall and F1 score, and a confusion matrix."
    )
    add_table(doc, ["Metric", "Meaning"], [
        ("Precision", "Among images predicted as a class, the fraction that truly belongs to that class."),
        ("Recall", "Among all validation images of a class, the fraction the model correctly finds."),
        ("F1 score", "A balance between precision and recall."),
        ("Support", "The number of validation images available for that class."),
        ("Confusion matrix", "A class-by-class count of correct predictions and confusions."),
    ], widths=[4.3, 12.7])

    add_heading(doc, "Single-image prediction", 2)
    doc.add_paragraph(
        "predict.py applies the same RGB conversion, resizing and tensor conversion used during validation. It loads the PyTorch "
        "checkpoint, applies softmax to the model logits and uses class_names.json to translate the largest score into a class name."
    )
    add_code(doc, "with torch.no_grad():\n    logits = model(image_tensor.unsqueeze(0))\n    probabilities = torch.softmax(logits, dim=1)[0]\nindex = int(probabilities.argmax())\npredicted_class = class_names[index]")
    doc.add_paragraph(
        "The printed confidence is a softmax score. It has not been calibrated against a large independent test set, so it should "
        "not be interpreted as a guaranteed real-world probability."
    )

    # 8
    add_heading(doc, "8 Deep learning fundamentals", 1)
    add_heading(doc, "Tensors", 2)
    doc.add_paragraph(
        "A tensor is a multidimensional numerical array. A scalar has zero dimensions, a vector has one, a matrix has two, and "
        "an image batch typically has four. The current PyTorch code stores image batches as batch, channels, height and width. "
        "TensorFlow commonly uses batch, height, width and channels."
    )
    add_table(doc, ["Concept", "TensorFlow convention", "PyTorch convention"], [
        ("One RGB image", "128 x 128 x 3", "3 x 128 x 128"),
        ("Batch of eight", "8 x 128 x 128 x 3", "8 x 3 x 128 x 128"),
        ("Data type", "tf.float32", "torch.float32"),
        ("Class labels", "Integer tensor", "Long integer tensor"),
    ], widths=[4.3, 6.2, 6.5])

    add_heading(doc, "Parameters and activations", 2)
    add_term(doc, "Parameter", "A learned weight or bias stored by a layer.")
    add_term(doc, "Activation", "The numerical output produced by a neuron or feature map.")
    add_term(doc, "ReLU", "The function max(0, x). Negative values become zero and positive values pass through.")
    add_term(doc, "Softmax", "Converts final class scores into positive values that sum to one.")

    add_heading(doc, "Forward pass and backpropagation", 2)
    doc.add_paragraph(
        "The forward pass calculates predictions from the current parameters. Backpropagation applies the chain rule from the loss "
        "back through every differentiable operation. Automatic differentiation in both TensorFlow and PyTorch records the graph "
        "of operations and calculates gradients for the trainable parameters. The optimizer then performs an update similar to:"
    )
    add_code(doc, "new_parameter = old_parameter - learning_rate * gradient")

    add_heading(doc, "Generalization, overfitting and underfitting", 2)
    add_term(doc, "Generalization", "Performance on genuinely new data from the same intended problem.")
    add_term(doc, "Overfitting", "Training performance improves while validation performance stops improving or becomes worse.")
    add_term(doc, "Underfitting", "Both training and validation performance remain weak because the model or training process has not learned enough.")
    doc.add_paragraph(
        "Augmentation, dropout, source-group splitting and early stopping reduce overfitting risk. They cannot replace independent "
        "field data. Strong evaluation still requires different plants, locations, cameras and capture sessions."
    )

    # 9
    add_heading(doc, "9 TensorFlow and PyTorch comparison", 1)
    doc.add_paragraph(
        "The project currently uses PyTorch. TensorFlow and Keras are included here for comparison because both frameworks "
        "implement the same core ideas: tensors, differentiable layers, losses, gradients and optimizers."
    )
    add_table(doc, ["Task", "TensorFlow and Keras", "PyTorch"], [
        ("Define model", "keras.Sequential or subclass keras.Model", "Subclass torch.nn.Module"),
        ("Convolution", "keras.layers.Conv2D", "torch.nn.Conv2d"),
        ("Pooling", "keras.layers.MaxPooling2D", "torch.nn.MaxPool2d"),
        ("Dense layer", "keras.layers.Dense", "torch.nn.Linear"),
        ("Dropout", "keras.layers.Dropout", "torch.nn.Dropout"),
        ("Loss", "sparse_categorical_crossentropy", "torch.nn.CrossEntropyLoss"),
        ("Optimizer", "keras.optimizers.Adam", "torch.optim.Adam"),
        ("Training", "model.fit", "Explicit Python batch loop"),
        ("Evaluation mode", "Handled by model.fit and predict", "model.eval with torch.no_grad"),
        ("Save model", ".keras model package", "Usually a .pth state dictionary"),
    ], widths=[3.5, 6.8, 6.7], font_size=8.1)

    add_heading(doc, "Previous TensorFlow model structure", 2)
    add_code(doc, "model = keras.Sequential([\n    keras.Input(shape=(128, 128, 3)),\n    keras.layers.RandomFlip(\"horizontal\"),\n    keras.layers.Conv2D(8, 3, activation=\"relu\", padding=\"same\"),\n    keras.layers.MaxPooling2D(),\n    keras.layers.Conv2D(16, 3, activation=\"relu\", padding=\"same\"),\n    keras.layers.MaxPooling2D(),\n    keras.layers.Conv2D(32, 3, activation=\"relu\", padding=\"same\"),\n    keras.layers.MaxPooling2D(),\n    keras.layers.Flatten(),\n    keras.layers.Dense(32, activation=\"relu\"),\n    keras.layers.Dropout(0.3),\n    keras.layers.Dense(13, activation=\"softmax\"),\n])")

    add_heading(doc, "Current PyTorch model structure", 2)
    add_code(doc, "class PlantCNN(nn.Module):\n    def __init__(self, num_classes=13):\n        super().__init__()\n        self.features = nn.Sequential(\n            nn.Conv2d(3, 8, 3, padding=1), nn.ReLU(), nn.MaxPool2d(2),\n            nn.Conv2d(8, 16, 3, padding=1), nn.ReLU(), nn.MaxPool2d(2),\n            nn.Conv2d(16, 32, 3, padding=1), nn.ReLU(), nn.MaxPool2d(2),\n        )\n        self.classifier = nn.Sequential(\n            nn.Flatten(), nn.Linear(32 * 16 * 16, 32),\n            nn.ReLU(), nn.Dropout(0.3), nn.Linear(32, num_classes),\n        )\n\n    def forward(self, x):\n        return self.classifier(self.features(x))")
    doc.add_paragraph(
        "PyTorch CrossEntropyLoss expects raw logits, so the PyTorch model does not place softmax inside forward during training. "
        "The loss function applies the required log-softmax operation internally."
    )

    add_heading(doc, "Keras fit compared with the current PyTorch loop", 2)
    add_table(doc, ["Training step", "Keras behavior", "PyTorch code"], [
        ("Clear gradients", "Managed by fit", "optimizer.zero_grad()"),
        ("Forward pass", "Managed by fit", "logits = model(images)"),
        ("Calculate loss", "Managed by compile and fit", "loss = criterion(logits, labels)"),
        ("Backpropagate", "Managed by fit", "loss.backward()"),
        ("Update weights", "Managed by fit", "optimizer.step()"),
    ], widths=[4.0, 6.0, 7.0])

    # 10
    add_heading(doc, "10 Current results and limitations", 1)
    add_table(doc, ["Measure", "Result"], [
        ("Training split", "194 images"),
        ("Validation split", "48 images"),
        ("Best checkpoint", "Epoch 23"),
        ("Validation accuracy", "26 of 48, or 54.17%"),
        ("Validation loss", "1.4408"),
        ("Independent test accuracy", "Unavailable"),
    ], widths=[7.0, 10.0], font_size=9.2)
    doc.add_paragraph(
        "The larger validation split gives a more stable diagnostic measurement than the earlier 23-image split. The score still "
        "does not estimate deployment performance because each class has only one conservative collection group and several classes "
        "have very few independent original sources."
    )
    add_heading(doc, "What the report currently shows", 2)
    reset_numbering()
    for text in [
        "The six newly confirmed classes perform better than several older provisional classes.",
        "Some old classes have only one to four validation images, so their individual scores are unstable.",
        "Class imbalance encourages the model to favor larger classes.",
        "Generated images increase visual variation but do not create new biological specimens.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Most valuable next improvements", 2)
    reset_numbering()
    for text in [
        "Collect independent photographs for every class and reserve them as a true test set.",
        "Confirm the remaining provisional botanical labels before treating the model as a species classifier.",
        "Balance the number of independent specimens per class rather than only balancing file counts.",
        "Compare the small CNN with transfer learning from MobileNetV2 or EfficientNetB0.",
        "Calibrate confidence and inspect repeated errors after an independent test set exists.",
    ]:
        add_numbered(doc, text)

    # 11
    add_heading(doc, "11 Running the project", 1)
    add_table(doc, ["Action", "Command", "Purpose"], [
        ("Check data", r".\.venv\Scripts\python.exe train.py --check-data", "Validate the manifest and images without training."),
        ("Train", r".\.venv\Scripts\python.exe train.py --diagnostic", "Use the source-group 80/20 split and train for up to 25 epochs."),
        ("Change epoch limit", r".\.venv\Scripts\python.exe train.py --diagnostic --epochs 40", "Allow up to 40 epochs while retaining early stopping."),
        ("Evaluate", r".\.venv\Scripts\python.exe evaluate.py", "Evaluate the saved best checkpoint on the saved holdout."),
        ("Predict", r".\.venv\Scripts\python.exe predict.py \"path\to\new_plant.jpg\"", "Classify one new image."),
    ], widths=[3.2, 8.3, 5.5], font_size=7.8)

    add_heading(doc, "Reading the outputs", 2)
    add_table(doc, ["Output", "What to inspect"], [
        ("results/accuracy.png", "Whether training and validation accuracy improve together."),
        ("results/loss.png", "Whether validation loss improves, plateaus or rises while training loss falls."),
        ("results/classification_report.txt", "Per-class precision, recall, F1 score and support."),
        ("results/confusion_matrix.png", "Which plant classes the model repeatedly confuses."),
        ("models/split.json", "The exact files assigned to training and validation."),
    ], widths=[6.2, 10.8])

    # 12
    add_heading(doc, "12 Glossary and study path", 1)
    glossary = [
        ("CNN", "A neural network designed to learn spatial image features with convolutional filters."),
        ("Feature map", "The spatial output created by a convolutional filter."),
        ("Kernel or filter", "A small learned grid that detects a local visual pattern."),
        ("Logit", "A raw class score before softmax."),
        ("Inference", "Using a trained model to predict without updating its parameters."),
        ("Checkpoint", "A saved model state selected during training."),
        ("Hyperparameter", "A user-selected setting such as learning rate, batch size or epoch limit."),
        ("Class imbalance", "A dataset condition where some classes have substantially more examples than others."),
        ("Data leakage", "Information from validation or test data unintentionally entering training."),
        ("Transfer learning", "Starting from a network pretrained on a large dataset and adapting it to a new task."),
    ]
    add_table(doc, ["Term", "Meaning"], glossary, widths=[4.4, 12.6])

    add_heading(doc, "Recommended learning sequence", 2)
    reset_numbering()
    for text in [
        "Run train.py --check-data and connect each printed check to load_dataset.",
        "Trace one image through RGB conversion, resizing, normalization and batching.",
        "Draw the tensor shape after each convolution and pooling layer.",
        "Review accuracy.png and loss.png alongside the concepts of underfitting and overfitting.",
        "Read the confusion matrix and identify which classes need more independent data.",
        "Compare the explicit PyTorch training loop with Keras model.fit to understand both APIs.",
        "After the baseline is understood, compare it with a transfer-learning model.",
    ]:
        add_numbered(doc, text)

    add_heading(doc, "Final understanding", 2)
    doc.add_paragraph(
        "The project is a complete baseline pipeline: it validates traceable image data, prevents direct source leakage, trains a "
        "small PyTorch CNN, selects a checkpoint by validation loss, evaluates class-level behavior and predicts new images. "
        "The code exposes the forward pass, gradient update and evaluation mode explicitly. TensorFlow can automate more of this "
        "loop through Keras model.fit, while the central "
        "deep learning ideas remain the same in both frameworks."
    )

    # Document properties and final consistency.
    doc.core_properties.title = "Plant Classification Training Guide"
    doc.core_properties.subject = "Project flow and PyTorch fundamentals with TensorFlow comparison"
    doc.core_properties.author = "Plant Classification Project"
    doc.core_properties.keywords = "plant classification, CNN, TensorFlow, PyTorch, deep learning"

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if not run.font.name:
                set_repeat_font(run)

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
